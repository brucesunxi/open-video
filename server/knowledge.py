import io
import re
import zipfile
import logging
from collections import Counter
import jieba
from rank_bm25 import BM25Okapi
from pypdf import PdfReader
from docx import Document
from .store import ident

jieba.setLogLevel(logging.WARNING)
STOP = set('的 了 呢 吗 呀 啊 和 是 在 有 请 什么 怎么 为什么 如何 可以 一个 一下 老师 我 你 它 这个 那个 告诉 介绍 讲解 说说'.split())


def tokens(text):
    return [s.lower() for s in jieba.lcut(text) if s.strip() and re.search(r'[\w\u4e00-\u9fff]', s) and s not in STOP]


def parse_document(name: str, content: bytes):
    ext = name.rsplit('.', 1)[-1].lower()
    if ext in ('txt', 'md'):
        return [(1, content.decode('utf-8-sig'))]
    if ext == 'pdf':
        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 300:
            raise ValueError('首版每份 PDF 最多 300 页，请拆分上传。')
        return [(i + 1, p.extract_text() or '') for i, p in enumerate(reader.pages)]
    if ext == 'docx':
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise ValueError('DOCX 解压体积过大，请拆分文档。')
        doc = Document(io.BytesIO(content))
        # DOCX has no stable page numbers: use logical section position instead.
        texts = [p.text for p in doc.paragraphs]
        texts.extend(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
        return [(None, '\n\n'.join(texts))]
    raise ValueError('知识文档支持 UTF-8 TXT、Markdown、DOCX 和文字 PDF。')


def chunk_pages(pages):
    chunks = []
    for page, text in pages:
        paragraphs = re.split(r'\n\s*\n|\n(?=#{1,6}\s)', text)
        for paragraph in paragraphs:
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            # Keep complete short paragraphs; split long text at sentence boundaries.
            parts = re.split(r'(?<=[。！？；.!?])\s*', paragraph) if len(paragraph) > 500 else [paragraph]
            buffer = ''
            for part in parts:
                if len(buffer) + len(part) > 650 and buffer:
                    chunks.append({'id': ident('chunk'), 'page': page, 'text': buffer})
                    buffer = ''
                # Hard limit pathological no-punctuation paragraphs.
                for offset in range(0, len(part), 650):
                    piece = part[offset:offset + 650]
                    if len(buffer) + len(piece) > 650:
                        chunks.append({'id': ident('chunk'), 'page': page, 'text': buffer})
                        buffer = ''
                    buffer += piece
            if buffer:
                chunks.append({'id': ident('chunk'), 'page': page, 'text': buffer})
    for i, chunk in enumerate(chunks):
        chunk['section'] = i + 1
    return chunks


def retrieve(question, documents, limit=3):
    query = tokens(question)
    if not query:
        return []
    candidates = []
    for doc in documents:
        if not doc.get('approved'):
            continue
        for chunk in doc['chunks']:
            candidates.append({**chunk, 'document_id': doc['id'], 'title': doc['title'], 'version': doc['version']})
    if not candidates:
        return []
    corpus = [tokens(c['text']) for c in candidates]
    scores = BM25Okapi(corpus).get_scores(query)
    qset = set(query)
    ranked = []
    for c, words, score in zip(candidates, corpus, scores):
        overlap = qset & set(words)
        # Baseline keyword gate. Not a semantic guarantee; demo returns source text only.
        meaningful = {t for t in overlap if len(t) > 1 or t.isdigit()}
        if not meaningful or len(overlap) / len(qset) < 0.35:
            continue
        c['score'] = round(float(score) + len(overlap) * 2, 3)
        ranked.append(c)
    return sorted(ranked, key=lambda x: x['score'], reverse=True)[:limit]
